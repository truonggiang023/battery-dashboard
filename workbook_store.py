"""Validated, versioned Excel storage for operational project records.

Workbooks contain input values, never synthetic results. Writes use a cell lock,
optimistic revisions, atomic replacement and immutable previous-version backups.
"""
from pathlib import Path
from contextlib import contextmanager
from contextvars import ContextVar
from datetime import date, datetime, timezone
import fcntl
import hashlib
import io
import json
import os
import tempfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import storage
from domain import CELLS, MODULES

# name: (type, required, allowed values). IDs must be unique within each sheet.
def text(required=False, choices=()): return ('text',required,choices)
def number(required=False): return ('number',required,())
def day(required=False): return ('date',required,())
DATASET=ContextVar('workbook_dataset',default='live')

@contextmanager
def scope(dataset):
    if dataset not in ('live','example'):raise ValueError('Dataset không hợp lệ.')
    token=DATASET.set(dataset)
    try:yield
    finally:DATASET.reset(token)

def data_root():
    return storage.ROOT/'example_data' if DATASET.get()=='example' else storage.ROOT

STATUS=('Open','In progress','Completed','Closed','Blocked','Cancelled')
APPROVAL=('Draft','Approved','Rejected')
SCHEMAS={
 'project': {'filename':'Project Information & KPI.xlsx','sheets':{
  'Project Information':{'project_id':text(True),'project_name':text(True),'cell_type':text(),'start_date':day(),'phase':text(),'owner':text(),'currency':text()},
  'Milestones':{'milestone_id':text(True),'name':text(True),'weight':number(True),'completion':number(True),'baseline':day(),'forecast':day(),'actual':day(),'status':text(True,('Planned','In progress','Completed','Approved','Blocked','Cancelled')),'pic':text(True),'depends_on':text()},
  'Budget':{'item_id':text(True),'description':text(True),'approved':number(),'actual':number(),'committed':number(),'forecast':number(),'currency':text(True),'pic':text(True)},
  'Issues':{'issue_id':text(True),'title':text(True),'severity':text(True,('Low','Medium','High','Critical')),'status':text(True,STATUS),'due_date':day(),'pic':text(True),'action':text(),'evidence':text()},
  'Tasks':{'task_id':text(True),'title':text(True),'workstream':text(),'priority':text(False,('Low','Medium','High','Critical')),'pic':text(True),'planned_date':day(),'due_date':day(),'status':text(True,STATUS),'metric':text(False,('energy','capacity','cycle','cost_prototype','cost_sop')),'milestone_id':text(),'issue_id':text(),'doe_id':text(),'build_lot_id':text(),'cell_id':text(),'description':text()},
  'PIC':{'activity_id':text(True),'department':text(),'role':text(),'pic':text(True),'backup':text(),'deadline':day(),'reporting_to':text()},
  'Weekly Report':{'record_id':text(True),'reporting_week':day(True),'task_id':text(),'status':text(False,STATUS),'evidence':text(),'next_week_plan':text(),'pic':text(True)},
  'DOE Runs':{'doe_id':text(True),'name':text(True),'status':text(True,('Planned','In progress','Completed','Approved','Cancelled')),'pic':text(True),'planned_date':day(),'completed_date':day(),'build_lot_id':text()},
 }},
 'targets':{'filename':'Target vs Actual.xlsx','sheets':{
  'Targets':{'metric_id':text(True,('energy','capacity','cycle','cost_prototype','cost_sop')),'target':number(True),'unit':text(True),'condition':text(),'doe_id':text(),'c_rate':number(),'temperature_c':number(),'retention_threshold':number(),'pic':text(True)},
  'Measurements':{'measurement_id':text(True),'metric_id':text(True,('energy','capacity')),'value':number(),'unit':text(),'discharge_energy_wh':number(),'cell_mass_kg':number(),'condition':text(),'c_rate':number(),'temperature_c':number(),'doe_id':text(),'build_lot_id':text(),'cell_id':text(),'test_date':day(True),'approval':text(True,APPROVAL),'revision':text(True),'pic':text(True)},
  'Cycle Prediction':{'prediction_id':text(True),'forecast_cycles':number(True),'retention_threshold':number(True),'model_revision':text(True),'validated':text(True,('Yes','No')),'condition':text(),'doe_id':text(),'build_lot_id':text(),'cell_id':text(),'prediction_date':day(True),'approval':text(True,APPROVAL),'pic':text(True)},
  'Cost Details':{'line_id':text(True),'cost_type':text(True,('Prototype','SOP')),'item':text(True),'quantity':number(True),'unit_price':number(True),'process_cost':number(),'overhead':number(),'currency':text(True),'approval':text(True,APPROVAL),'revision':text(True),'updated_at':day(True),'pic':text(True)},
 }}
}
# One workbook per functional item; child content is represented by sheets.
BASE=SCHEMAS
from catalog import CATALOG
import re

def sheet_name(label):
    return re.sub(r'[\\/*?:\[\]:]', '-', label)[:31].strip() or 'Information'

def generic_fields(information=''):
    import unicodedata
    fields={'entry_id':text(True)}
    for label in information.split(';'):
        label=unicodedata.normalize('NFKD',label).encode('ascii','ignore').decode().lower()
        key=re.sub(r'[^a-z0-9]+','_',label).strip('_')[:60]
        if key and key not in fields:fields[key]=text()
    for key,value in {'pic':text(True),'status':text(),'revision':text(),'approval':text(False,APPROVAL),'doe_id':text(),'build_lot_id':text(),'cell_id':text(),'notes':text()}.items():
        fields[key]=value
    return fields

SCHEMAS={
 'project':dict(filename='Project Information & KPI.xlsx',module=0,section='Project Information & KPI',sheets={
     'Project Information':BASE['project']['sheets']['Project Information'],
     'Budget':BASE['project']['sheets']['Budget']}),
 'targets':dict(filename='Target vs Actual.xlsx',module=0,section='Target vs Actual',sheets={
     'Targets':BASE['targets']['sheets']['Targets'],
     'Energy density':BASE['targets']['sheets']['Measurements'],
     'Cell capacity':BASE['targets']['sheets']['Measurements'],
     'Cycle life predicted':BASE['targets']['sheets']['Cycle Prediction'],
     'Cell cost Prototype':BASE['targets']['sheets']['Cost Details'],
     'Cell cost SOP':BASE['targets']['sheets']['Cost Details']}),
 'timeline':dict(filename='Master Timeline.xlsx',module=0,section='Master Timeline',sheets={'Project Milestones':BASE['project']['sheets']['Milestones']}),
 'issues':dict(filename='Rủi ro & Sự vụ.xlsx',module=0,section='Rủi ro & Sự vụ',sheets={'Issue Register':BASE['project']['sheets']['Issues']}),
 'execution':dict(filename='Execution Control.xlsx',module=0,section='Execution Control',sheets={
     'Task Master List':BASE['project']['sheets']['Tasks'],
     'Action Tracking':dict(BASE['project']['sheets']['Tasks']),
     'Deliverables':{'deliverable_id':text(True),'title':text(True),'revision':text(True),'approver':text(),'status':text(False,APPROVAL),'pic':text(True),'due_date':day(),'document_ref':text()}}),
 'pic':dict(filename='Cấu trúc & PIC.xlsx',module=0,section='Cấu trúc & PIC',sheets={
     'Organization Structure':{'role_id':text(True),'department':text(True),'role':text(True),'pic':text(True),'reporting_to':text()},
     'PIC Assignment':BASE['project']['sheets']['PIC']}),
 'weekly':dict(filename='Báo cáo tuần.xlsx',module=0,section='Báo cáo tuần',sheets={name:dict(BASE['project']['sheets']['Weekly Report']) for name in ['Weekly Progress','Completed Activities','Next Week Plan','Open Issues']}),
 'doe_runs':dict(filename='DOE Run Tracking.xlsx',module=3,section='DOE Run Tracking',sheets={'DOE Runs':BASE['project']['sheets']['DOE Runs']}),
}
SCHEMAS['performance']=dict(filename='Electrical Performance.xlsx',module=6,section='Electrical Performance',sheets={
 'Performance Curves':{'record_id':text(True),'series':text(True),'cycle_number':number(),'retention_pct':number(),'voltage':number(),'discharged_capacity_ah':number(),'time_s':number(),'current_a':number(),'doe_id':text(),'cell_id':text(),'approval':text(True,APPROVAL),'pic':text(True)}})
SCHEMAS['digital_twin']=dict(filename='Cycler - 3D Cell Digital Twin.xlsx',module=8,section='Cycler · 3D Cell Digital Twin',sheets={
 'Cell Mapping':{'cell_id':text(True),'model':text(True),'form_factor':text(True,('Cylindrical','Prismatic')),'diameter_mm':number(),'height_mm':number(True),'width_mm':number(),'thickness_mm':number(),'model_revision':text(True),'geometry_ref':text(),'device':text(),'channel':text(),'pic':text(True)},
 'Cell State':{'record_id':text(True),'cell_id':text(True),'timestamp':text(True),'soc_pct':number(),'voltage_v':number(),'current_a':number(),'temperature_c':number(),'cycle_number':number(),'capacity_ah':number(),'retention_pct':number(),'status':text(False,('Idle','Charging','Discharging','Rest','Fault')),'pic':text(True)}})
for module,sections in CATALOG['sections'].items():
    for section in sections:
        if any(book['module']==int(module) and book['section']==section for book in SCHEMAS.values()):continue
        content=[r['content'] for r in CATALOG['records'] if r['module']==int(module) and r['section']==section]
        names=list(dict.fromkeys(sheet_name(label) for label in content)) or ['Information']
        key='function_'+str(module)+'_'+hashlib.sha256(section.encode()).hexdigest()[:10]
        SCHEMAS[key]=dict(filename=re.sub(r'[\\/*?:\[\]:]','-',section)+'.xlsx',module=int(module),section=section,sheets={name:generic_fields(next((r['information'] for r in CATALOG['records'] if r['module']==int(module) and r['section']==section and sheet_name(r['content'])==name),'')) for name in names})

class DataError(ValueError): pass
class Conflict(DataError): pass

def path_for(cell,book):
    if cell not in CELLS or book not in SCHEMAS: raise DataError('Cell hoặc workbook không hợp lệ.')
    c=CELLS[cell]
    return data_root()/'projects'/c['project']/c['chemistry']/c['size']/MODULES[SCHEMAS[book]['module']][0]/SCHEMAS[book]['filename']

@contextmanager
def lock(cell):
    directory=data_root()/'workbook_locks'; directory.mkdir(parents=True,exist_ok=True)
    with (directory/f'{cell}.lock').open('a') as handle:
        fcntl.flock(handle,fcntl.LOCK_EX)
        try: yield
        finally: fcntl.flock(handle,fcntl.LOCK_UN)

def revision(content):return hashlib.sha256(content).hexdigest()

def build(book,sheets):
    wb=Workbook();wb.remove(wb.active)
    for name,fields in SCHEMAS[book]['sheets'].items():
        ws=wb.create_sheet(name);ws.append(list(fields));ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{ws.cell(1,len(fields)).column_letter}{max(2,len(sheets.get(name,[]))+1)}'
        for cell in ws[1]:
            cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='1B354A');cell.alignment=Alignment(wrap_text=True)
            ws.column_dimensions[cell.column_letter].width=min(36,max(18,len(str(cell.value))+4))
        for row in sheets.get(name,[]):
            ws.append([row.get(key) for key in fields])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        for col,(_key,(kind,required,choices)) in enumerate(fields.items(),1):
            if choices:
                dv=DataValidation(type='list',formula1='"'+','.join(choices)+'"',allow_blank=not required);dv.errorTitle='Giá trị không hợp lệ';dv.error='Chọn một giá trị trong danh sách.';dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(f'{ws.cell(1,col).column_letter}2:{ws.cell(1,col).column_letter}5001')
        ws.sheet_properties.pageSetUpPr.fitToPage=True
    output=io.BytesIO();wb.save(output);wb.close();return output.getvalue()

def initialize(cells=None):
    for cell in cells or CELLS:
        with lock(cell):
            for book in SCHEMAS:
                path=path_for(cell,book);path.parent.mkdir(parents=True,exist_ok=True)
                if not path.exists():
                    with path.open('xb') as handle: handle.write(build(book,{}))

def normalize(book,sheet,rows):
    if not isinstance(rows,list) or len(rows)>5000:raise DataError('Sheet chỉ hỗ trợ tối đa 5000 dòng.')
    fields=SCHEMAS[book]['sheets'][sheet];result=[];ids=set()
    for index,row in enumerate(rows,2):
        if not isinstance(row,dict) or set(row)-set(fields):raise DataError(f'{sheet}, dòng {index}: sai tên cột.')
        if not any(v not in ('',None) for v in row.values()):continue
        record={}
        for key,(kind,required,choices) in fields.items():
            value=row.get(key)
            if value is None or (isinstance(value,str) and not value.strip()):
                if required:raise DataError(f'{sheet}, dòng {index}: thiếu {key}.')
                record[key]=None;continue
            if kind=='number':
                import math
                if isinstance(value,bool):raise DataError(f'{sheet}, dòng {index}: {key} phải là số.')
                try:value=float(value)
                except (TypeError,ValueError):raise DataError(f'{sheet}, dòng {index}: {key} phải là số.')
                if not math.isfinite(value) or (value<0 and key not in ('temperature_c','current_a')):raise DataError(f'{sheet}, dòng {index}: {key} không hợp lệ.')
            elif kind=='date':
                try:value=(value.date() if isinstance(value,datetime) else value) if isinstance(value,(date,datetime)) else date.fromisoformat(str(value))
                except ValueError:raise DataError(f'{sheet}, dòng {index}: {key} dùng YYYY-MM-DD.')
                value=value.isoformat()
            else:
                value=str(value).strip()
                if len(value)>2000 or value.startswith('='):raise DataError(f'{sheet}, dòng {index}: {key} quá dài hoặc chứa công thức Excel. Hãy nhập giá trị.')
                if choices and value not in choices:raise DataError(f'{sheet}, dòng {index}: {key} phải thuộc {", ".join(choices)}.')
            record[key]=value
        ident=record[next(iter(fields))]
        if ident in ids:raise DataError(f'{sheet}, dòng {index}: ID bị trùng ({ident}).')
        ids.add(ident)
        if record.get('c_rate') is not None and record['c_rate']<=0:raise DataError('c_rate phải >0.')
        if record.get('planned_date') and record.get('due_date') and record['due_date']<record['planned_date']:raise DataError('due_date không được trước planned_date.')
        if sheet=='DOE Runs' and record['status'] in ('Completed','Approved') and not record['completed_date']:raise DataError('DOE Completed/Approved cần completed_date.')
        if sheet=='Project Milestones':
            if record['completion']>100:raise DataError('Milestones: completion phải từ 0 đến 100.')
            if record['status'] in ('Completed','Approved') and (record['completion']!=100 or not record['actual']):raise DataError('Gate Completed/Approved phải có completion=100 và actual date.')
        if sheet in ('Targets','Cycle life predicted') and record.get('retention_threshold') is not None and not 0<record['retention_threshold']<=100:raise DataError('retention_threshold phải >0 và <=100 (%).')
        if sheet in ('Energy density','Cell capacity'):
            if record['metric_id']!=('energy' if sheet=='Energy density' else 'capacity'):raise DataError('metric_id không khớp sheet.')
            if record['metric_id']=='energy' and (record['discharge_energy_wh'] is None or not record['cell_mass_kg'] or record.get('unit') not in (None,'Wh/kg')):raise DataError('Energy cần discharge_energy_wh, cell_mass_kg > 0; unit=Wh/kg hoặc để trống.')
            if record['metric_id']=='capacity' and (record['value'] is None or record['unit']!='Ah' or record['c_rate'] is None or record['temperature_c'] is None):raise DataError('Capacity cần value, unit=Ah, c_rate và temperature_c.')
        if sheet=='Targets':
            expected={'energy':'Wh/kg','capacity':'Ah','cycle':'cycles'}.get(record['metric_id'])
            if expected and record['unit']!=expected:raise DataError(f"{record['metric_id']}: unit phải là {expected}.")
            if record['target']<=0:raise DataError('Target phải >0.')
            if record['metric_id'].startswith('cost_') and not record['unit'].endswith('/cell'):raise DataError('Cost unit phải có dạng USD/cell hoặc VND/cell.')
            if record['metric_id']=='cycle' and record['retention_threshold'] is None:raise DataError('Target cycle cần retention_threshold (%).')
        if sheet in ('Cell cost Prototype','Cell cost SOP') and record['cost_type']!=('Prototype' if sheet=='Cell cost Prototype' else 'SOP'):raise DataError('cost_type không khớp sheet.')
        if sheet=='Performance Curves' and record.get('retention_pct') is not None and record['retention_pct']>100:raise DataError('retention_pct phải từ 0 đến 100.')
        if sheet=='Cell Mapping':
            if not record['height_mm']:raise DataError('Chiều cao cell phải >0.')
            if record['form_factor']=='Cylindrical' and not record['diameter_mm']:raise DataError('Cell trụ cần diameter_mm >0.')
            if record['form_factor']=='Prismatic' and (not record['width_mm'] or not record['thickness_mm']):raise DataError('Cell prismatic cần width_mm và thickness_mm >0.')
        if sheet=='Cell State':
            try:
                stamp=datetime.fromisoformat(record['timestamp'].replace('Z','+00:00'))
                if stamp.tzinfo is None:raise ValueError()
            except ValueError:raise DataError('timestamp phải là ISO-8601 có múi giờ, ví dụ 2026-10-20T09:00:00+07:00.')
            for field in ('soc_pct','retention_pct'):
                if record[field] is not None and record[field]>100:raise DataError(f'{field} phải từ 0 đến 100.')
        result.append(record)
    if sheet=='Cell State' and len({(r['cell_id'],datetime.fromisoformat(r['timestamp'].replace('Z','+00:00')).timestamp()) for r in result})!=len(result):raise DataError('Cell State: trùng cell_id và timestamp.')
    if sheet=='Cell Mapping':
        occupied=[(r['device'],r['channel']) for r in result if r['device'] and r['channel']]
        if len(set(occupied))!=len(occupied):raise DataError('Cell Mapping: một device/channel không được gán hai cell cùng lúc.')
    if sheet=='Project Information' and len(result)>1:raise DataError('Project Information chỉ có 1 dòng cho mỗi cell.')
    return result

def read(cell,book):
    path=path_for(cell,book)
    if not path.exists():raise DataError('Workbook chưa tồn tại. Khởi tạo workbook trước.')
    content=path.read_bytes()
    if len(content)>10_000_000:raise DataError('Workbook vượt giới hạn 10 MB.')
    try:
        import zipfile
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist())>50_000_000:raise DataError('Workbook giải nén vượt giới hạn 50 MB.')
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
    except DataError:raise
    except Exception:raise DataError(f'{SCHEMAS[book]["filename"]}: không đọc được Excel; kiểm tra file đã lưu hoàn chỉnh.')
    sheets={}
    try:
        if set(wb.sheetnames)!=set(SCHEMAS[book]['sheets']):raise DataError('Workbook có sheet thiếu hoặc ngoài cấu hình. Không lưu để tránh mất dữ liệu; giữ nguyên các sheet của file mẫu.')
        for name,fields in SCHEMAS[book]['sheets'].items():
            if name not in wb.sheetnames:raise DataError(f'Thiếu sheet {name}.')
            ws=wb[name]
            if ws.max_row>5001:raise DataError(f'{name}: quá 5000 dòng.')
            values=ws.iter_rows(values_only=True);headers=list(next(values,()))
            if headers!=list(fields):raise DataError(f'{name}: sai tên hoặc thứ tự cột. Hãy dùng file mẫu của ứng dụng.')
            sheets[name]=normalize(book,name,[dict(zip(headers,row)) for row in values])
    finally:wb.close()
    return {'revision':revision(content),'sheets':sheets,'modified_at':datetime.fromtimestamp(path.stat().st_mtime,timezone.utc).isoformat()}

def save_sheet(cell,book,sheet,rows,expected_revision,actor="local"):
    if book not in SCHEMAS or sheet not in SCHEMAS[book]['sheets']:raise DataError('Workbook hoặc sheet không hợp lệ.')
    normalized=normalize(book,sheet,rows)
    with lock(cell):
        current=read(cell,book)
        if not expected_revision or current['revision']!=expected_revision:raise Conflict('File đã thay đổi. Tải lại sheet trước khi lưu để tránh ghi đè dữ liệu của người khác.')
        path=path_for(cell,book);previous=path.read_bytes()
        if revision(previous)!=current['revision']:raise Conflict('File Excel vừa thay đổi bên ngoài. Tải lại trước khi lưu.')
        backup=data_root()/'workbook_history'/cell/book;backup.mkdir(parents=True,exist_ok=True)
        old=backup/f'{current["revision"]}.xlsx'
        if not old.exists():
            with old.open('xb') as handle:handle.write(previous);handle.flush();os.fsync(handle.fileno())
        current['sheets'][sheet]=normalized
        if book=='digital_twin':
            mapped={r['cell_id'] for r in current['sheets']['Cell Mapping']}
            if any(r['cell_id'] not in mapped for r in current['sheets']['Cell State']):raise DataError('Cell State chứa Cell ID chưa có trong Cell Mapping.')
        content=build(book,current['sheets'])
        fd,temporary=tempfile.mkstemp(dir=path.parent,suffix='.xlsx')
        try:
            with os.fdopen(fd,'wb') as handle:handle.write(content);handle.flush();os.fsync(handle.fileno())
            if revision(path.read_bytes())!=current['revision']:raise Conflict('File Excel vừa thay đổi. Dữ liệu chưa được ghi đè; hãy tải lại file.')
            os.replace(temporary,path)
            descriptor=os.open(path.parent,os.O_RDONLY)
            try:os.fsync(descriptor)
            finally:os.close(descriptor)
        finally:
            if os.path.exists(temporary):os.unlink(temporary)
        with (backup/'audit.jsonl').open('a') as audit:
            audit.write(json.dumps(dict(timestamp=datetime.now(timezone.utc).isoformat(),actor=actor,book=book,sheet=sheet,previous=current['revision'],revision=revision(content),rows=len(normalized)),ensure_ascii=False)+'\n')
            audit.flush();os.fsync(audit.fileno())
        return read(cell,book)
