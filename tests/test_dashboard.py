import unittest, tempfile, io, os
from pathlib import Path
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
import storage
import workbook_store as store
import live_data, example_project
from app import app

CELL='CYL-LFP-40135';EXAMPLE=example_project.EXAMPLE_CELL
HEADERS={'X-Requested-With':'BatteryDashboard'}

class DashboardTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.original=storage.ROOT
        storage.ROOT=Path(self.tmp.name);store.DATASET.set('live');storage.initialize();store.initialize();example_project.initialize();self.client=app.test_client()
    def tearDown(self):storage.ROOT=self.original;store.DATASET.set('live');self.tmp.cleanup()
    def book(self,key='targets',dataset='example',cell=EXAMPLE):
        return self.client.get(f'/api/cells/{cell}/workbooks/{key}?dataset={dataset}').json
    def save(self,key,sheet,rows,rev,dataset='example',cell=EXAMPLE):
        return self.client.put(f'/api/cells/{cell}/workbooks/{key}/sheets/{sheet}?dataset={dataset}',json=dict(rows=rows,revision=rev),headers=HEADERS)
    def summary(self,dataset='example',cell=EXAMPLE):return self.client.get(f'/api/cells/{cell}/summary?dataset={dataset}').json
    def test_live_empty_and_example_separation(self):
        actual=self.summary('live',CELL);example=self.summary()
        self.assertIsNone(actual['progress']);self.assertIsNone(actual['metrics']['capacity']['actual']);self.assertEqual(actual['color'],'Grey');self.assertFalse(actual['demo'])
        self.assertEqual(example['progress'],68);self.assertEqual(example['milestones']['completed'],2);self.assertEqual(example['milestones']['total'],7);self.assertEqual(example['issues']['open'],4);self.assertTrue(example['demo'])
        self.assertEqual(example['budget']['remaining'],16_000_000_000)
        self.assertEqual(example['metrics']['energy']['actual'],178)
        self.assertEqual(example['metrics']['cost_prototype']['actual'],22.6)
        self.assertEqual(example['metrics']['cost_sop']['actual'],17.8)
    def test_approved_inputs_drive_kpi_and_preserve_previous_workbook(self):
        book=self.book();energy=book['sheets']['Energy density'];energy[0]['discharge_energy_wh']=380
        response=self.save('targets','Energy density',energy,book['revision']);self.assertEqual(response.status_code,200)
        self.assertEqual(self.summary()['metrics']['energy']['actual'],200)
        self.assertIsNone(self.summary('live',CELL)['metrics']['energy']['actual'])
        backup=storage.ROOT/'example_data'/'workbook_history'/EXAMPLE/'targets'/f"{book['revision']}.xlsx"
        with backup.open('rb') as handle:
            wb=load_workbook(handle,read_only=True);self.assertEqual(list(wb['Energy density'].values)[1][4],338.2);wb.close()
        book=self.book();energy=book['sheets']['Energy density'];energy[0]['approval']='Draft'
        self.assertEqual(self.save('targets','Energy density',energy,book['revision']).status_code,200)
        self.assertEqual(self.summary()['metrics']['energy']['actual'],175)
        self.assertEqual(self.summary()['metrics']['energy']['date'],'2026-10-13')
        book=self.book();energy=book['sheets']['Energy density']
        for record in energy:record['approval']='Draft'
        self.assertEqual(self.save('targets','Energy density',energy,book['revision']).status_code,200)
        self.assertIsNone(self.summary()['metrics']['energy']['actual'])
    def test_revision_conflict_does_not_overwrite(self):
        book=self.book();rows=book['sheets']['Targets'];rows[0]['target']=190
        self.assertEqual(self.save('targets','Targets',rows,book['revision']).status_code,200)
        rows[0]['target']=200
        self.assertEqual(self.save('targets','Targets',rows,book['revision']).status_code,409)
        self.assertEqual(self.book()['sheets']['Targets'][0]['target'],190)
    def test_external_excel_update_is_detected_without_restart(self):
        with store.scope('example'):path=store.path_for(EXAMPLE,'targets')
        wb=load_workbook(path);ws=wb['Cell capacity'];headers=[c.value for c in ws[1]];ws.cell(2,headers.index('value')+1,104.5);wb.save(path);wb.close()
        self.assertEqual(self.summary()['metrics']['capacity']['actual'],104.5)
    def test_mismatched_conditions_are_not_aggregated(self):
        book=self.book();rows=book['sheets']['Cell capacity'];other=dict(rows[0],measurement_id='Q002',condition='1C / 25C',c_rate=1,value=90);rows.append(other)
        self.assertEqual(self.save('targets','Cell capacity',rows,book['revision']).status_code,200)
        self.assertEqual(self.summary()['metrics']['capacity']['actual'],103.2)
        book=self.book();targets=book['sheets']['Targets'];targets[1]['condition']=None;targets[1]['c_rate']=None
        self.assertEqual(self.save('targets','Targets',targets,book['revision']).status_code,200)
        result=self.summary()['metrics']['capacity'];self.assertIsNone(result['actual']);self.assertTrue(result['warning'])
    def test_validation_rejects_bad_units_ids_and_false_completed_gate(self):
        book=self.book();rows=book['sheets']['Targets'];rows[0]['unit']='g/cm3'
        self.assertEqual(self.save('targets','Targets',rows,book['revision']).status_code,422)
        book=self.book();rows=book['sheets']['Targets'];rows.append(dict(rows[0]))
        self.assertEqual(self.save('targets','Targets',rows,book['revision']).status_code,422)
        book=self.book('timeline');rows=book['sheets']['Project Milestones'];rows[2]['status']='Completed'
        self.assertEqual(self.save('timeline','Project Milestones',rows,book['revision']).status_code,422)
        self.assertEqual(self.book('timeline')['revision'],book['revision'])
    def test_concurrent_writers_only_one_succeeds(self):
        store.DATASET.set('live');current=store.read(CELL,'project')
        row=dict(project_id='P01',project_name='Actual project',cell_type='40135',owner='PMO',currency='VND')
        def update(name):
            try:return store.save_sheet(CELL,'project','Project Information',[dict(row,project_name=name)],current['revision'])['revision']
            except store.Conflict:return 'conflict'
        with ThreadPoolExecutor(max_workers=2) as pool:outcomes=list(pool.map(update,['Team A','Team B']))
        self.assertEqual(outcomes.count('conflict'),1)
        self.assertEqual(len(store.read(CELL,'project')['sheets']['Project Information']),1)
    def test_digital_twin_mapping_and_state_link(self):
        data=self.client.get(f'/api/cells/{EXAMPLE}/digital-twin?dataset=example').json
        self.assertTrue(data['mapping']);self.assertTrue(data['states']);self.assertFalse(data['live_connection'])
        book=self.book('digital_twin');rows=book['sheets']['Cell State'];rows[0]['cell_id']='MISSING'
        self.assertEqual(self.save('digital_twin','Cell State',rows,book['revision']).status_code,422)
        self.assertEqual(self.book('digital_twin')['revision'],book['revision'])
    def test_auth_and_write_request_guard(self):
        with patch.dict(os.environ,{'DASHBOARD_USERNAME':'test-user','DASHBOARD_PASSWORD':'test-only-pass'}):
            self.assertEqual(self.client.get('/api/hierarchy').status_code,401)
            self.assertEqual(self.client.get('/api/hierarchy',auth=('test-user','test-only-pass')).status_code,200)
        response=self.client.put(f'/api/cells/{CELL}/workbooks/project/sheets/Project Information',json={})
        self.assertEqual(response.status_code,403)
        self.assertEqual(self.client.get('/api/cells/unknown/summary').status_code,404)
        self.assertEqual(self.client.get('/api/hierarchy?dataset=../escape').status_code,400)
    def test_workbook_download_and_task_links_use_actual_records(self):
        for group in ('progress','budget','milestones','issues','energy','capacity','cycle','cost_prototype','cost_sop'):
            result=self.client.get(f'/api/cells/{EXAMPLE}/tasks/{group}?dataset=example')
            self.assertEqual(result.status_code,200);self.assertTrue(result.json['rows'])
            empty=self.client.get(f'/api/cells/{CELL}/tasks/{group}')
            self.assertEqual(empty.json['rows'],[])
        with self.client.get(f'/api/cells/{EXAMPLE}/workbooks/targets/download?dataset=example') as response:
            self.assertEqual(response.status_code,200)
            wb=load_workbook(io.BytesIO(response.data),read_only=True);self.assertIn('Energy density',wb.sheetnames);wb.close()
        import zipfile
        with self.client.get(f'/api/cells/{EXAMPLE}/excel-package?dataset=example') as response:
            self.assertEqual(response.status_code,200)
            with zipfile.ZipFile(io.BytesIO(response.data)) as package:
                self.assertEqual(sum(name.endswith('.xlsx') for name in package.namelist()),len(store.SCHEMAS))
        materials=next(key for key,value in store.SCHEMAS.items() if value['module']==2 and value['section']=='Cathode')
        fields=self.book(materials)['fields']['Active material - Grade']
        self.assertTrue({'material_id','supplier','composition','lot','pic'}<=set(f['name'] for f in fields))
    def test_unknown_sheet_is_preserved_by_rejecting_rewrite(self):
        with store.scope('example'):path=store.path_for(EXAMPLE,'project')
        wb=load_workbook(path);wb.create_sheet('Team Notes')['A1']='Preserve me';wb.save(path);wb.close();before=path.read_bytes()
        response=self.client.get(f'/api/cells/{EXAMPLE}/workbooks/project?dataset=example')
        self.assertEqual(response.status_code,422);self.assertEqual(path.read_bytes(),before)
    def test_original_digital_twin_model_and_live_control_isolation(self):
        from integrations.digital_twin.rack import Rack
        rack=Rack();self.assertEqual(len(rack.cells),12);self.assertFalse(rack.running)
        rack.pending_fault={'cell':1,'kind':'overheat'};rack.running=True;rack.tick(.1)
        self.assertTrue(rack.safety_stop);self.assertFalse(rack.running);self.assertTrue(all(c['current_a']==0 for c in rack.cells))
        response=self.client.post('/api/cycler/control',json={'action':'start'},headers=HEADERS)
        self.assertEqual(response.status_code,403)
        response=self.client.post('/api/cycler/control?dataset=example',json={'action':'speed','speed':60},headers=HEADERS)
        self.assertEqual(response.status_code,200)
        response=self.client.post('/api/cycler/control?dataset=example',json={'action':'speed','speed':99},headers=HEADERS)
        self.assertEqual(response.status_code,409)

if __name__=='__main__':unittest.main()
